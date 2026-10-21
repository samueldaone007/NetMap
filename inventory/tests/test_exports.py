from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from inventory.models import Connection, Device, IPAddress, Location, Outage, Port, Subnet


class ExportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("keeper", password="test-password-123")
        self.client.force_login(self.user)
        self.location = Location.objects.create(site="Riverside", building="Annex", rack="R1")
        self.subnet = Subnet.objects.create(cidr="10.20.30.0/24", name="Office", vlan_id=30)
        self.device = Device.objects.create(
            name="core-sw-01",
            device_type=Device.Type.SWITCH,
            vendor="Cisco",
            model="C9200",
            serial_number="SN-01",
            management_ip="10.20.30.2",
            location=self.location,
            purpose="=SUM(1,2)",
            notes="Main distribution switch",
        )
        self.port_a = Port.objects.create(device=self.device, label="Gi0/1", status=Port.Status.UP)
        self.server = Device.objects.create(name="srv-01", device_type=Device.Type.SERVER)
        self.port_b = Port.objects.create(
            device=self.server, label="eth0", status=Port.Status.IN_USE
        )
        Connection.objects.create(
            port_a=self.port_a, port_b=self.port_b, cable_label="C-001"
        )
        IPAddress.objects.create(
            address="10.20.30.20",
            subnet=self.subnet,
            device=self.server,
            hostname="srv-01.office",
            status=IPAddress.Status.USED,
        )
        Outage.objects.create(
            device=self.server,
            started=timezone.now() - timedelta(hours=2),
            cause="Power feed",
        )

    def test_export_index_shows_live_manifest(self):
        response = self.client.get(reverse("export_index"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Take the atlas with you.")
        self.assertContains(response, "Inventory workbook")

    def test_workbook_contains_all_sheets_and_formula_safe_text(self):
        response = self.client.get(reverse("export_workbook"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment;", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(
            workbook.sheetnames,
            [
                "Summary",
                "Devices",
                "Ports",
                "Cables",
                "Subnets",
                "IP Addresses",
                "Locations",
                "Outages",
                "Change Log",
                "Imports",
            ],
        )
        devices = workbook["Devices"]
        headers = [cell.value for cell in devices[1]]
        purpose_column = headers.index("Purpose") + 1
        name_column = headers.index("Device") + 1
        row = next(
            row
            for row in devices.iter_rows(min_row=2)
            if row[name_column - 1].value == "core-sw-01"
        )
        self.assertEqual(row[purpose_column - 1].value, "'=SUM(1,2)")
        cables = workbook["Cables"]
        self.assertEqual(cables.max_row, 2)

    def test_pdf_export_is_a_real_pdf(self):
        response = self.client.get(reverse("export_pdf"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        self.assertGreater(len(response.content), 1000)

    def test_export_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("export_workbook"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response["Location"])
