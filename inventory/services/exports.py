"""Workbook and printable PDF exports."""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    HRFlowable,
    LongTable,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    TableStyle,
)

from ..models import (
    ChangeLog,
    Connection,
    Device,
    ImportBatch,
    IPAddress,
    Location,
    Outage,
    Port,
    Subnet,
)

ILLEGAL_XLSX = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
MAX_CELL = 32767

PAPER = "F3EDE2"
PAPER_LIGHT = "FAF6EE"
INK = "2B2622"
MUTED = "746A5F"
RULE = "C9BCA9"
MUSTARD = "E0A526"
PLUM = "7B3F6E"
MOSS = "5E7F3A"
COBALT = "2F5DA8"


def inventory_counts():
    return {
        "locations": Location.objects.count(),
        "subnets": Subnet.objects.count(),
        "devices": Device.objects.count(),
        "ips": IPAddress.objects.count(),
        "ports": Port.objects.count(),
        "connections": Connection.objects.count(),
        "outages": Outage.objects.filter(ended__isnull=True).count(),
        "placeholders": Device.objects.filter(needs_identifying=True).count(),
    }


def _now_label():
    return timezone.localtime().strftime("%Y-%m-%d %H:%M")


def _excel_safe(value):
    """Keep exported text inert and within Excel's cell limits."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.replace(tzinfo=None)
    if isinstance(value, (int, float, date)):
        return value
    text = ILLEGAL_XLSX.sub("", str(value))
    if text.startswith(FORMULA_PREFIXES):
        text = "'" + text
    if len(text) > MAX_CELL:
        text = text[: MAX_CELL - 16] + "… [truncated]"
    return text


def _connections():
    return list(
        Connection.objects.select_related(
            "port_a__device", "port_b__device"
        ).order_by("port_a__device__name", "port_a__label")
    )


def _connected_index(connections):
    connected = defaultdict(list)
    for connection in connections:
        a = f"{connection.port_a.device.name}:{connection.port_a.label}"
        if connection.port_b_id:
            b = f"{connection.port_b.device.name}:{connection.port_b.label}"
        else:
            b = "Unknown"
        connected[connection.port_a_id].append(b)
        if connection.port_b_id:
            connected[connection.port_b_id].append(a)
    return connected


def _port_rows(connections):
    connected = _connected_index(connections)
    rows = []
    ports = Port.objects.select_related("device").order_by("device__name", "label")
    for port in ports:
        used = (
            port.status in {Port.Status.UP, Port.Status.IN_USE}
            or bool(connected.get(port.pk))
        )
        rows.append(
            [
                port.device.name,
                port.label,
                port.get_port_type_display(),
                port.speed,
                port.poe,
                port.access_vlan,
                port.tagged_vlans,
                port.connector,
                port.fibre_type,
                port.provider,
                port.circuit_id,
                port.get_status_display(),
                "Yes" if used else "No",
                ", ".join(connected.get(port.pk, [])),
            ]
        )
    return rows


def _device_rows(connections):
    connected = _connected_index(connections)
    rows = []
    devices = (
        Device.objects.select_related("location")
        .prefetch_related("ports", "ip_addresses", "outages")
        .order_by("name")
    )
    for device in devices:
        ports = list(device.ports.all())
        used = sum(
            1
            for port in ports
            if port.status in {Port.Status.UP, Port.Status.IN_USE}
            or bool(connected.get(port.pk))
        )
        open_outage = next(
            (outage.cause for outage in device.outages.all() if outage.ended is None),
            "",
        )
        rows.append(
            [
                device.pk,
                device.name,
                device.get_device_type_display(),
                device.vendor,
                device.model,
                device.serial_number or "",
                device.asset_id or "",
                device.hostname,
                device.os_version,
                device.firmware,
                device.management_ip or "",
                str(device.location) if device.location else "",
                device.owner,
                device.purpose,
                device.purchase_date,
                device.warranty_expiry,
                device.get_status_display(),
                device.needs_identifying,
                device.get_source_display(),
                device.last_verified,
                used,
                len(ports),
                ", ".join(
                    sorted(ip.address for ip in device.ip_addresses.all())
                ),
                open_outage,
                device.notes,
            ]
        )
    return rows


def build_workbook():
    now = timezone.localtime()
    counts = inventory_counts()
    connections = _connections()

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "NetMap"
    workbook.properties.title = "NetMap inventory"
    workbook.properties.description = "Local network inventory export"

    summary = workbook.create_sheet("Summary")
    summary.sheet_properties.tabColor = MUSTARD
    summary["A1"] = "NETMAP"
    summary["A1"].font = Font(name="Arial", size=24, bold=True, color=INK)
    summary["A2"] = "Inventory workbook"
    summary["A2"].font = Font(name="Arial", size=14, color=PLUM)
    summary["A3"] = f"Generated {now.strftime('%Y-%m-%d %H:%M')} · local single-user inventory"
    summary["A3"].font = Font(name="Arial", size=9, color=MUTED)
    summary["A5"] = "Section"
    summary["B5"] = "Records"
    for cell in (summary["A5"], summary["B5"]):
        cell.font = Font(name="Arial", size=9, bold=True, color=PAPER_LIGHT)
        cell.fill = PatternFill("solid", fgColor=INK)
    labels = [
        ("Sites", counts["locations"]),
        ("Subnets", counts["subnets"]),
        ("Devices", counts["devices"]),
        ("IP addresses", counts["ips"]),
        ("Ports", counts["ports"]),
        ("Cables", counts["connections"]),
        ("Open outages", counts["outages"]),
        ("Placeholders to identify", counts["placeholders"]),
    ]
    for offset, (label, value) in enumerate(labels, start=6):
        summary.cell(offset, 1, label).font = Font(name="Arial", size=10)
        summary.cell(offset, 2, value).font = Font(name="Arial", size=10, bold=True)
        summary.cell(offset, 2).alignment = Alignment(horizontal="right")
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 14
    note_row = 6 + len(labels) + 2
    summary.cell(note_row, 1, "Text that begins with =, +, - or @ is written as plain text.")
    summary.cell(note_row, 1).font = Font(name="Arial", size=8, italic=True, color=MUTED)
    summary.freeze_panes = "A6"

    sheets = [
        (
            "Devices",
            PLUM,
            [
                "ID",
                "Device",
                "Type",
                "Vendor",
                "Model",
                "Serial number",
                "Asset ID",
                "Hostname",
                "OS / version",
                "Firmware",
                "Management IP",
                "Location",
                "Owner",
                "Purpose",
                "Purchase date",
                "Warranty expiry",
                "Status",
                "Needs identifying",
                "Source",
                "Last verified",
                "Ports used",
                "Ports total",
                "IP addresses",
                "Open outage",
                "Notes",
            ],
            _device_rows(connections),
        ),
        (
            "Ports",
            COBALT,
            [
                "Device",
                "Label",
                "Type",
                "Speed",
                "PoE",
                "Access VLAN",
                "Tagged VLANs",
                "Connector",
                "Fibre type",
                "Provider",
                "Circuit ID",
                "Status",
                "Used",
                "Connected to",
            ],
            _port_rows(connections),
        ),
        (
            "Cables",
            MUSTARD,
            [
                "Port A device",
                "Port A",
                "Port B device",
                "Port B",
                "Cable label",
                "Notes",
                "Source",
                "State",
            ],
            [
                [
                    connection.port_a.device.name,
                    connection.port_a.label,
                    connection.port_b.device.name if connection.port_b_id else "",
                    connection.port_b.label if connection.port_b_id else "",
                    connection.cable_label,
                    connection.notes,
                    connection.get_source_display(),
                    "Unknown far end" if connection.port_b_id is None else "Traced",
                ]
                for connection in connections
            ],
        ),
        (
            "Subnets",
            MOSS,
            [
                "CIDR",
                "VLAN",
                "Name",
                "Gateway",
                "Purpose",
                "Usable hosts",
                "Used",
                "Recorded",
                "Created",
            ],
            [
                [
                    subnet.cidr,
                    subnet.vlan_id,
                    subnet.name,
                    subnet.gateway or "",
                    subnet.purpose,
                    subnet.host_count,
                    subnet.used_count,
                    subnet.ips.count(),
                    subnet.created_at,
                ]
                for subnet in Subnet.objects.prefetch_related("ips").order_by("cidr")
            ],
        ),
        (
            "IP Addresses",
            COBALT,
            [
                "Address",
                "Subnet",
                "Subnet name",
                "Device",
                "Hostname",
                "Status",
                "Source",
                "Last verified",
            ],
            [
                [
                    ip.address,
                    ip.subnet.cidr,
                    ip.subnet.name,
                    ip.device.name if ip.device_id else "",
                    ip.hostname,
                    ip.get_status_display(),
                    ip.get_source_display(),
                    ip.last_verified,
                ]
                for ip in IPAddress.objects.select_related("subnet", "device").order_by(
                    "subnet__cidr", "address"
                )
            ],
        ),
        (
            "Locations",
            MUSTARD,
            ["Site", "Building", "Room", "Rack", "Rack unit", "Devices", "Notes"],
            [
                [
                    location.site,
                    location.building,
                    location.room,
                    location.rack,
                    location.rack_unit,
                    location.devices.count(),
                    location.notes,
                ]
                for location in Location.objects.prefetch_related("devices").order_by(
                    "site", "building", "room", "rack", "rack_unit"
                )
            ],
        ),
        (
            "Outages",
            PLUM,
            ["Started", "Ended", "Open", "Subject", "Cause", "Notes"],
            [
                [
                    outage.started,
                    outage.ended,
                    outage.is_open,
                    str(outage.subject),
                    outage.cause,
                    outage.notes,
                ]
                for outage in Outage.objects.select_related(
                    "device", "connection__port_a__device", "connection__port_b__device"
                ).order_by("-started")
            ],
        ),
        (
            "Change Log",
            MUTED,
            ["Timestamp", "Object type", "Object", "Action", "Field", "Old", "New", "User"],
            [
                [
                    entry.timestamp,
                    entry.object_type,
                    entry.object_repr,
                    entry.get_action_display(),
                    entry.field,
                    entry.old_value,
                    entry.new_value,
                    entry.user.username if entry.user_id else "system",
                ]
                for entry in ChangeLog.objects.select_related("user").order_by(
                    "-timestamp", "-pk"
                )
            ],
        ),
        (
            "Imports",
            INK,
            ["File", "Sheet", "Contents", "Status", "Rows", "Created", "Imported", "Undone"],
            [
                [
                    batch.original_name,
                    batch.sheet_name,
                    batch.get_target_type_display(),
                    batch.get_status_display(),
                    batch.row_count,
                    batch.created_at,
                    batch.imported_at,
                    batch.undone_at,
                ]
                for batch in ImportBatch.objects.all()
            ],
        ),
    ]

    thin = Side(style="thin", color=RULE)
    border = Border(bottom=thin)
    for title, tab_color, headers, rows in sheets:
        sheet = workbook.create_sheet(title)
        sheet.sheet_properties.tabColor = tab_color
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(name="Arial", size=9, bold=True, color=PAPER_LIGHT)
            cell.fill = PatternFill("solid", fgColor=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 27
        for row in rows:
            sheet.append([_excel_safe(value) for value in row])
        for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
            for cell in row:
                cell.font = Font(name="Arial", size=9, color=INK)
                cell.border = border
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
        widths = []
        for index, header in enumerate(headers, start=1):
            longest = len(str(header))
            for cell in sheet[get_column_letter(index)][: min(sheet.max_row, 160)]:
                value = cell.value
                if isinstance(value, (date, datetime)):
                    value = value.strftime("%Y-%m-%d")
                if value is not None:
                    longest = max(longest, len(str(value)))
            widths.append(min(max(longest + 2, 10), 46))
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        if sheet.max_row > 1:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
            )
        sheet.sheet_view.showGridLines = False

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

PDF_INK = colors.HexColor("#2B2622")
PDF_MUTED = colors.HexColor("#746A5F")
PDF_RULE = colors.HexColor("#C9BCA9")
PDF_PAPER = colors.HexColor("#F3EDE2")
PDF_PAPER_LIGHT = colors.HexColor("#FAF6EE")
PDF_MUSTARD = colors.HexColor("#E0A526")
PDF_PLUM = colors.HexColor("#7B3F6E")
PDF_MOSS = colors.HexColor("#5E7F3A")
PDF_COBALT = colors.HexColor("#2F5DA8")


def _pdf_styles():
    base = getSampleStyleSheet()
    return {
        "eyebrow": ParagraphStyle(
            "NetMapEyebrow",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7.5,
            leading=10,
            textColor=PDF_PLUM,
            spaceAfter=4,
        ),
        "title": ParagraphStyle(
            "NetMapTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=26,
            leading=28,
            textColor=PDF_INK,
            alignment=TA_LEFT,
            spaceAfter=4,
        ),
        "lede": ParagraphStyle(
            "NetMapLede",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=PDF_MUTED,
            spaceAfter=10,
        ),
        "section": ParagraphStyle(
            "NetMapSection",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=14,
            leading=16,
            textColor=PDF_INK,
            spaceBefore=14,
            spaceAfter=6,
        ),
        "small": ParagraphStyle(
            "NetMapSmall",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=7,
            leading=9,
            textColor=PDF_MUTED,
        ),
        "cell": ParagraphStyle(
            "NetMapCell",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=6.6,
            leading=8.2,
            textColor=PDF_INK,
        ),
        "cell_mono": ParagraphStyle(
            "NetMapCellMono",
            parent=base["Normal"],
            fontName="Courier",
            fontSize=6.3,
            leading=8,
            textColor=PDF_INK,
        ),
        "head": ParagraphStyle(
            "NetMapHead",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=6.2,
            leading=7.5,
            textColor=PDF_PAPER_LIGHT,
        ),
    }


def _pdf_safe(value, style):
    text = "" if value is None else str(value)
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    return Paragraph(text, style)


def _pdf_date(value):
    if not value:
        return ""
    if isinstance(value, datetime):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")
    return value.strftime("%Y-%m-%d")


def _pdf_table(headers, rows, widths, styles, mono_columns=()):
    data = [[_pdf_safe(header, styles["head"]) for header in headers]]
    for row in rows:
        cells = []
        for index, value in enumerate(row):
            style = styles["cell_mono"] if index in mono_columns else styles["cell"]
            if isinstance(value, datetime):
                value = _pdf_date(value)
            elif isinstance(value, date):
                value = _pdf_date(value)
            cells.append(_pdf_safe(value, style))
        data.append(cells)
    table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), PDF_INK),
                ("TEXTCOLOR", (0, 0), (-1, 0), PDF_PAPER_LIGHT),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LINEBELOW", (0, 0), (-1, -1), 0.35, PDF_RULE),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [PDF_PAPER_LIGHT, colors.white]),
            ]
        )
    )
    return table


def _page_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(PDF_RULE)
    canvas.setLineWidth(0.5)
    canvas.line(18 * mm, 15 * mm, A4[0] - 18 * mm, 15 * mm)
    canvas.setFont("Helvetica", 6.5)
    canvas.setFillColor(PDF_MUTED)
    canvas.drawString(18 * mm, 10 * mm, "NetMap · local inventory · 127.0.0.1")
    canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf():
    counts = inventory_counts()
    connections = _connections()
    now = _now_label()
    styles = _pdf_styles()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=21 * mm,
        title="NetMap inventory",
        author="NetMap",
        subject="Local network inventory",
    )
    story = []
    story.append(Paragraph("NETMAP · INVENTORY DOCKET", styles["eyebrow"]))
    story.append(Paragraph("The network, on paper.", styles["title"]))
    story.append(
        Paragraph(
            f"Generated {now} from the local NetMap inventory. "
            "This is a human-kept record, not a live scan.",
            styles["lede"],
        )
    )
    story.append(
        HRFlowable(
            width="100%",
            thickness=2,
            color=PDF_INK,
            spaceBefore=0,
            spaceAfter=9,
        )
    )
    summary_data = [
        [
            _pdf_safe("Sites", styles["small"]),
            _pdf_safe(counts["locations"], styles["section"]),
            _pdf_safe("Subnets", styles["small"]),
            _pdf_safe(counts["subnets"], styles["section"]),
            _pdf_safe("Devices", styles["small"]),
            _pdf_safe(counts["devices"], styles["section"]),
        ],
        [
            _pdf_safe("IP addresses", styles["small"]),
            _pdf_safe(counts["ips"], styles["section"]),
            _pdf_safe("Ports", styles["small"]),
            _pdf_safe(counts["ports"], styles["section"]),
            _pdf_safe("Cables", styles["small"]),
            _pdf_safe(counts["connections"], styles["section"]),
        ],
    ]
    summary = LongTable(
        summary_data,
        colWidths=[22 * mm, 18 * mm, 22 * mm, 18 * mm, 22 * mm, 18 * mm],
        hAlign="LEFT",
    )
    summary.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 2),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LINEBELOW", (0, 1), (-1, 1), 0.5, PDF_RULE),
            ]
        )
    )
    story.append(summary)
    if counts["placeholders"]:
        story.append(
            Paragraph(
                f"{counts['placeholders']} placeholder record(s) still need identifying.",
                styles["lede"],
            )
        )

    story.append(Paragraph("Devices", styles["section"]))
    device_rows = _device_rows(connections)
    device_pdf_rows = [
        [
            row[1],
            row[2],
            row[11],
            row[10],
            f"{row[20]}/{row[21]}",
            row[16],
            row[19] or "",
        ]
        for row in device_rows
    ]
    if device_pdf_rows:
        story.append(
            _pdf_table(
                ["Device", "Type", "Location", "Mgmt IP", "Ports", "Status", "Verified"],
                device_pdf_rows,
                [44 * mm, 18 * mm, 39 * mm, 26 * mm, 12 * mm, 18 * mm, 17 * mm],
                styles,
                mono_columns=(3,),
            )
        )
    else:
        story.append(Paragraph("No devices recorded.", styles["lede"]))

    story.append(Paragraph("Subnets", styles["section"]))
    subnet_rows = [
        [
            subnet.cidr,
            subnet.vlan_id or "",
            subnet.name,
            subnet.gateway or "",
            f"{subnet.used_count} used / {subnet.ips.count()} recorded",
        ]
        for subnet in Subnet.objects.prefetch_related("ips").order_by("cidr")
    ]
    if subnet_rows:
        story.append(
            _pdf_table(
                ["CIDR", "VLAN", "Name", "Gateway", "Addresses"],
                subnet_rows,
                [34 * mm, 12 * mm, 46 * mm, 32 * mm, 50 * mm],
                styles,
                mono_columns=(0, 3),
            )
        )
    else:
        story.append(Paragraph("No subnets recorded.", styles["lede"]))

    story.append(Paragraph("IP address ledger", styles["section"]))
    ip_rows = [
        [
            ip.address,
            ip.subnet.cidr,
            ip.device.name if ip.device_id else "",
            ip.hostname,
            ip.get_status_display(),
            _pdf_date(ip.last_verified),
        ]
        for ip in IPAddress.objects.select_related("subnet", "device").order_by(
            "subnet__cidr", "address"
        )
    ]
    if ip_rows:
        story.append(
            _pdf_table(
                ["Address", "Subnet", "Device", "Hostname", "Status", "Verified"],
                ip_rows,
                [30 * mm, 28 * mm, 40 * mm, 38 * mm, 16 * mm, 22 * mm],
                styles,
                mono_columns=(0, 1),
            )
        )
    else:
        story.append(Paragraph("No addresses recorded.", styles["lede"]))

    story.append(Paragraph("Cables", styles["section"]))
    cable_rows = [
        [
            f"{connection.port_a.device.name}:{connection.port_a.label}",
            (
                f"{connection.port_b.device.name}:{connection.port_b.label}"
                if connection.port_b_id
                else "Unknown"
            ),
            connection.cable_label,
            connection.get_source_display(),
        ]
        for connection in connections
    ]
    if cable_rows:
        story.append(
            _pdf_table(
                ["Port A", "Port B", "Cable", "Source"],
                cable_rows,
                [63 * mm, 63 * mm, 30 * mm, 18 * mm],
                styles,
                mono_columns=(0, 1),
            )
        )
    else:
        story.append(Paragraph("No cables recorded.", styles["lede"]))

    story.append(Spacer(1, 8))
    story.append(
        Paragraph(
            "NetMap keeps the inventory local. Port usage is derived from port status and "
            "recorded connections; it is never stored as a separate total.",
            styles["small"],
        )
    )
    doc.build(story, onFirstPage=_page_footer, onLaterPages=_page_footer)
    return buffer.getvalue()
