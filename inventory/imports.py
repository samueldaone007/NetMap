"""Spreadsheet import: reading, cleanup, mapping, preview, commit and undo.

The importer is deliberately honest about damaged spreadsheet cells. A date in
a port-label column is an error, not a silent conversion. Free-text references
become flagged placeholders rather than failed rows.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal

import ipaddress
from django.apps import apps
from django.conf import settings
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from .models import Connection, Device, IPAddress, ImportBatch, Location, Port, Source
from .services.ports import expand_port_labels

EMPTY_VALUES = {"", "nil", "n/a", "na", "-", "none", "null"}
UNKNOWN_VALUES = {"unknown", "?", "??", "not known", "unpatched"}
TEXT_KINDS = {"text", "choice", "ip", "ref"}
DATE_FORMATS = (
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%d.%m.%Y",
    "%d-%m-%Y",
    "%Y/%m/%d",
    "%d %b %Y",
    "%b %d %Y",
    "%d %B %Y",
    "%B %d %Y",
)
MAX_ROWS = 5000
MAX_COLUMNS = 100


class ImportFileError(Exception):
    pass


class ImportCommitError(Exception):
    pass


class ImportUndoError(Exception):
    pass


@dataclass(frozen=True)
class ImportField:
    name: str
    label: str
    kind: str = "text"
    choices: tuple = ()

    @property
    def help(self):
        if self.kind == "ref":
            return "Existing name, serial or asset ID; free text becomes a placeholder."
        if self.kind == "date":
            return "Date cell or YYYY-MM-DD."
        return ""


def _choices(model_field):
    return tuple((str(value), str(label)) for value, label in model_field.choices)


DEVICE_FIELDS = [
    ImportField("name", "Device name"),
    ImportField("serial_number", "Serial number"),
    ImportField("asset_id", "Asset ID"),
    ImportField("device_type", "Device type", "choice", _choices(Device._meta.get_field("device_type"))),
    ImportField("vendor", "Vendor"),
    ImportField("model", "Model"),
    ImportField("hostname", "Hostname"),
    ImportField("os_version", "OS / version"),
    ImportField("firmware", "Firmware"),
    ImportField("management_ip", "Management IP", "ip"),
    ImportField("owner", "Owner"),
    ImportField("purpose", "Purpose"),
    ImportField("purchase_date", "Purchase date", "date"),
    ImportField("warranty_expiry", "Warranty expiry", "date"),
    ImportField("status", "Status", "choice", _choices(Device._meta.get_field("status"))),
    ImportField("notes", "Notes"),
    ImportField("last_verified", "Last verified", "date"),
    ImportField("location", "Location (free text)"),
    ImportField("site", "Location: site"),
    ImportField("building", "Location: building"),
    ImportField("room", "Location: room"),
    ImportField("rack", "Location: rack"),
    ImportField("rack_unit", "Location: rack unit (U)", "int"),
]

PORT_FIELDS = [
    ImportField("device", "Device", "ref"),
    ImportField("label", "Port label / range"),
    ImportField("port_type", "Port type", "choice", _choices(Port._meta.get_field("port_type"))),
    ImportField("speed", "Speed"),
    ImportField("poe", "PoE", "bool"),
    ImportField("access_vlan", "Access VLAN", "int"),
    ImportField("tagged_vlans", "Tagged VLANs"),
    ImportField("connector", "Connector"),
    ImportField("fibre_type", "Fibre type"),
    ImportField("provider", "Provider"),
    ImportField("circuit_id", "Circuit ID"),
    ImportField("status", "Status", "choice", _choices(Port._meta.get_field("status"))),
]

CONNECTION_FIELDS = [
    ImportField("port_a", "Port A (device:port or free text)", "ref"),
    ImportField("port_a_device", "Port A device"),
    ImportField("port_a_label", "Port A label"),
    ImportField("port_b", "Port B (device:port or free text)", "ref"),
    ImportField("port_b_device", "Port B device"),
    ImportField("port_b_label", "Port B label"),
    ImportField("cable_label", "Cable label"),
    ImportField("notes", "Notes"),
]

TARGET_FIELDS = {
    "devices": DEVICE_FIELDS,
    "ports": PORT_FIELDS,
    "connections": CONNECTION_FIELDS,
}

TARGET_LABELS = {
    "devices": "Devices",
    "ports": "Ports",
    "connections": "Connections",
}


def fields_for(target: str) -> list[ImportField]:
    return TARGET_FIELDS.get(target, [])


def field_names_for(target: str) -> set[str]:
    return {spec.name for spec in fields_for(target)}


# ---------------------------------------------------------------------------
# File reading and upload validation
# ---------------------------------------------------------------------------

def safe_original_name(name: str) -> str:
    name = os.path.basename(name or "upload").strip()
    name = re.sub(r"[\x00-\x1f\x7f]", "", name)
    return name[:255] or "upload"


def validate_upload(upload) -> str:
    if not upload:
        raise ImportFileError("Choose an .xlsx or .csv file to upload.")
    name = safe_original_name(upload.name)
    extension = os.path.splitext(name)[1].lower()
    if extension not in {".xlsx", ".csv"}:
        raise ImportFileError("Only .xlsx and .csv files are accepted. Macros and .xls are not.")
    max_bytes = settings.NETMAP_MAX_UPLOAD_MB * 1024 * 1024
    if upload.size <= 0:
        raise ImportFileError("That file is empty.")
    if upload.size > max_bytes:
        raise ImportFileError(
            f"That file is larger than {settings.NETMAP_MAX_UPLOAD_MB} MB."
        )
    head = upload.read(8)
    upload.seek(0)
    if extension == ".xlsx":
        if not head.startswith(b"PK\x03\x04"):
            raise ImportFileError("That file is not a valid .xlsx workbook.")
    else:
        if b"\x00" in head:
            raise ImportFileError("That does not look like a text CSV file.")
    return name


def sheet_names(batch: ImportBatch) -> list[str]:
    path = batch.uploaded_file.path
    if path.lower().endswith(".csv"):
        return ["CSV"]
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        names = list(workbook.sheetnames)
        workbook.close()
    except Exception as exc:
        raise ImportFileError(f"Could not read this workbook: {exc}") from exc
    if not names:
        raise ImportFileError("This workbook has no sheets.")
    if len(names) > 50:
        raise ImportFileError("This workbook has too many sheets for a manual import.")
    return names


@dataclass
class SheetData:
    headers: list[str]
    rows: list[tuple[int, list]]
    column_count: int


def _trim_row(values):
    values = list(values)
    while values and (values[-1] is None or (isinstance(values[-1], str) and not values[-1].strip())):
        values.pop()
    return values


def _read_xlsx(path: str, sheet: str) -> list[tuple[int, list]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise ImportFileError(f"Could not open workbook: {exc}") from exc
    if sheet not in workbook.sheetnames:
        workbook.close()
        raise ImportFileError(f"Sheet {sheet!r} is not in this workbook.")
    worksheet = workbook[sheet]
    if worksheet.max_column > MAX_COLUMNS:
        workbook.close()
        raise ImportFileError(
            f"Sheet {sheet!r} has {worksheet.max_column} columns; the limit is {MAX_COLUMNS}."
        )
    rows = []
    for number, raw in enumerate(worksheet.iter_rows(values_only=True), start=1):
        values = _trim_row(raw)
        if not values:
            continue
        if len(values) > MAX_COLUMNS:
            workbook.close()
            raise ImportFileError(f"Row {number} has more than {MAX_COLUMNS} columns.")
        rows.append((number, values))
        if len(rows) > MAX_ROWS + 1:
            workbook.close()
            raise ImportFileError(f"This sheet has more than {MAX_ROWS} rows.")
    workbook.close()
    return rows


def _read_csv(path: str) -> list[tuple[int, list]]:
    raw = open(path, "rb").read()
    if b"\x00" in raw:
        raise ImportFileError("The CSV contains binary data.")
    text = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportFileError("Could not decode the CSV. Save it as UTF-8 and try again.")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    csv.field_size_limit(200_000)
    rows = []
    reader = csv.reader(io.StringIO(text), dialect)
    for number, raw_row in enumerate(reader, start=1):
        values = _trim_row(raw_row)
        if not values:
            continue
        if len(values) > MAX_COLUMNS:
            raise ImportFileError(f"Row {number} has more than {MAX_COLUMNS} columns.")
        rows.append((number, values))
        if len(rows) > MAX_ROWS + 1:
            raise ImportFileError(f"This CSV has more than {MAX_ROWS} rows.")
    return rows


def read_sheet(batch: ImportBatch, sheet: str, has_header: bool) -> SheetData:
    path = batch.uploaded_file.path
    rows = _read_csv(path) if path.lower().endswith(".csv") else _read_xlsx(path, sheet)
    if not rows:
        raise ImportFileError("That sheet is empty.")
    column_count = max(len(values) for _, values in rows)
    if has_header:
        header_values = rows[0][1]
        headers = []
        for index in range(column_count):
            value = header_values[index] if index < len(header_values) else ""
            label = str(value).strip() if value is not None else ""
            headers.append(label or f"Column {index + 1}")
        data = rows[1:]
    else:
        headers = [f"Column {index + 1}" for index in range(column_count)]
        data = rows
    if not data:
        raise ImportFileError("That sheet only has a header row.")
    return SheetData(headers=headers, rows=data, column_count=column_count)


# ---------------------------------------------------------------------------
# Cleanup and parsing
# ---------------------------------------------------------------------------

def normalize_blank(raw):
    if raw is None:
        return None
    if isinstance(raw, str):
        text = raw.replace("\u00a0", " ").strip()
        if text.lower() in EMPTY_VALUES:
            return None
        return text
    return raw


def display_cell(raw):
    if raw is None:
        return ""
    if isinstance(raw, datetime):
        return raw.strftime("%Y-%m-%d")
    if isinstance(raw, date):
        return raw.isoformat()
    return str(raw).strip()


def _damaged_text(raw) -> str | None:
    if isinstance(raw, datetime):
        return f"looks like a date ({raw.date().isoformat()}) where text was expected"
    if isinstance(raw, date):
        return f"looks like a date ({raw.isoformat()}) where text was expected"
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return "looks like an Excel number/serial where text was expected"
    return None


def parse_text(raw, label):
    raw = normalize_blank(raw)
    if raw is None:
        return None, None
    damage = _damaged_text(raw)
    if damage:
        return None, (
            f"{label} {damage}. Format the cell as text in Excel and re-upload; "
            "NetMap will not guess."
        )
    return str(raw).strip(), None


def parse_choice(raw, label, choices):
    text, issue = parse_text(raw, label)
    if issue or text is None:
        return None, issue
    for value, choice_label in choices:
        if text.casefold() in {value.casefold(), choice_label.casefold()}:
            return value, None
    valid = ", ".join(choice_label for _, choice_label in choices)
    return None, f"{label}: {text!r} is not one of {valid}."


def parse_int(raw, label, minimum=None, maximum=None):
    raw = normalize_blank(raw)
    if raw is None:
        return None, None
    if isinstance(raw, (datetime, date)):
        return None, f"{label} looks like a date; expected a whole number."
    if isinstance(raw, bool):
        return None, f"{label} must be a whole number."
    try:
        if isinstance(raw, float) and not raw.is_integer():
            raise ValueError
        number = int(raw)
    except (TypeError, ValueError):
        return None, f"{label}: {raw!r} is not a whole number."
    if minimum is not None and number < minimum:
        return None, f"{label} must be at least {minimum}."
    if maximum is not None and number > maximum:
        return None, f"{label} must be no more than {maximum}."
    return number, None


def parse_bool(raw, label):
    raw = normalize_blank(raw)
    if raw is None:
        return None, None
    if isinstance(raw, bool):
        return raw, None
    if isinstance(raw, (int, float)) and raw in (0, 1):
        return bool(raw), None
    if isinstance(raw, str):
        lowered = raw.casefold()
        if lowered in {"yes", "y", "true", "1", "poe", "on"}:
            return True, None
        if lowered in {"no", "n", "false", "0", "non-poe", "off"}:
            return False, None
    return None, f"{label}: {raw!r} is not yes or no."


def parse_date(raw, label):
    raw = normalize_blank(raw)
    if raw is None:
        return None, None
    if isinstance(raw, datetime):
        return raw.date(), None
    if isinstance(raw, date):
        return raw, None
    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
        return None, f"{label} looks like an Excel date serial; format it as a date and re-upload."
    text = str(raw).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date(), None
        except ValueError:
            continue
    return None, f"{label}: {text!r} is not a recognisable date."


def parse_ip(raw, label):
    text, issue = parse_text(raw, label)
    if issue or text is None:
        return None, issue
    try:
        return str(ipaddress.ip_address(text)), None
    except ValueError:
        return None, f"{label}: {text!r} is not a valid IP address."


def parse_fields(row, mapping, target):
    """Return (clean_values, issues)."""
    specs = {spec.name: spec for spec in fields_for(target)}
    values = {}
    issues = []
    for index, field_name in mapping.items():
        if not field_name:
            continue
        try:
            column = int(index)
        except (TypeError, ValueError):
            issues.append(f"Mapping has an invalid column index {index!r}.")
            continue
        if column >= len(row):
            continue
        spec = specs.get(field_name)
        if spec is None:
            issues.append(f"Column {index + 1} is mapped to unknown field {field_name!r}.")
            continue
        raw = row[column]
        if spec.kind == "text":
            value, issue = parse_text(raw, spec.label)
        elif spec.kind == "ref":
            value, issue = parse_text(raw, spec.label)
        elif spec.kind == "choice":
            value, issue = parse_choice(raw, spec.label, spec.choices)
        elif spec.kind == "int":
            minimum = 1 if field_name == "rack_unit" else None
            maximum = 99 if field_name == "rack_unit" else 4094 if field_name in {"access_vlan"} else None
            value, issue = parse_int(raw, spec.label, minimum, maximum)
        elif spec.kind == "date":
            value, issue = parse_date(raw, spec.label)
        elif spec.kind == "bool":
            value, issue = parse_bool(raw, spec.label)
        elif spec.kind == "ip":
            value, issue = parse_ip(raw, spec.label)
        else:
            value, issue = parse_text(raw, spec.label)
        if issue:
            issues.append(issue)
        elif value is not None:
            values[field_name] = value
    return values, issues


def _fingerprint(target, values):
    payload = json.dumps(
        {key: str(value) for key, value in sorted(values.items()) if value not in (None, "")},
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(f"{target}:{payload}".encode("utf-8")).hexdigest()


def _safe_reference_fingerprint(reference):
    return hashlib.sha256(f"reference:{reference.strip().lower()}".encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# Preview data
# ---------------------------------------------------------------------------

@dataclass
class PreviewRow:
    number: int
    action: str
    subject: str
    details: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    expanded: str = ""


@dataclass
class ImportResult:
    rows: list[PreviewRow]
    counts: dict
    created_count: int = 0
    updated_count: int = 0
    placeholder_count: int = 0
    global_errors: list[str] = field(default_factory=list)
    total_rows: int = 0
    sheet: str = ""
    target: str = ""

    @property
    def has_errors(self):
        return bool(self.global_errors) or any(row.errors for row in self.rows) or any(
            row.action == "error" for row in self.rows
        )

    @property
    def error_count(self):
        return len(self.global_errors) + sum(1 for row in self.rows if row.errors or row.action == "error")


# ---------------------------------------------------------------------------
# Engine
# ---------------------------------------------------------------------------

class ImportEngine:
    def __init__(self, batch: ImportBatch):
        self.batch = batch
        self.target = batch.target_type
        self.mapping = batch.mapping or {}
        self.rows: list[PreviewRow] = []
        self.global_errors: list[str] = []
        self.created: list = []
        self.created_ids: set[int] = set()
        self.updated: dict = {}
        self.snapshots: dict = {}
        self.placeholder_count = 0

        self.devices_by_name = {d.name.casefold(): d for d in Device.objects.all()}
        self.devices_by_serial = {
            d.serial_number.casefold(): d for d in Device.objects.all() if d.serial_number
        }
        self.devices_by_asset = {
            d.asset_id.casefold(): d for d in Device.objects.all() if d.asset_id
        }
        self.devices_by_fingerprint = {
            d.import_fingerprint: d for d in Device.objects.all() if d.import_fingerprint
        }
        self.ports_by_key = {
            (port.device_id, port.label.casefold()): port
            for port in Port.objects.select_related("device").all()
        }
        self.connection_cache = {}
        for connection in Connection.objects.all():
            key = self._connection_key(connection.port_a_id, connection.port_b_id, raw=True)
            self.connection_cache[key] = connection
        self.location_cache = {}
        for location in Location.objects.all():
            self.location_cache[self._location_key_from_obj(location)] = location
            self.location_cache.setdefault(("str", str(location).casefold()), location)

    # -- object tracking -------------------------------------------------
    def _is_created(self, obj):
        return id(obj) in self.created_ids

    def _mark_created(self, obj):
        if not self._is_created(obj):
            self.created.append(obj)
            self.created_ids.add(id(obj))

    def _snapshot(self, obj):
        key = (obj._meta.label_lower, obj.pk)
        if key not in self.snapshots:
            self.snapshots[key] = serialize_instance(obj)
            self.updated[key] = obj

    def _assign(self, obj, fields: dict):
        if not self._is_created(obj):
            self._snapshot(obj)
        for name, value in fields.items():
            setattr(obj, name, value)

    # -- locations -------------------------------------------------------
    @staticmethod
    def _location_key_from_obj(location):
        return (
            "components",
            (location.site or "").casefold(),
            (location.building or "").casefold(),
            (location.room or "").casefold(),
            (location.rack or "").casefold(),
            location.rack_unit,
        )

    def _location_from_values(self, values):
        component_names = ("site", "building", "room", "rack", "rack_unit")
        if any(name in values for name in component_names):
            site = values.get("site") or "Unplaced site"
            key = (
                "components",
                site.casefold(),
                str(values.get("building") or "").casefold(),
                str(values.get("room") or "").casefold(),
                str(values.get("rack") or "").casefold(),
                values.get("rack_unit"),
            )
            if key in self.location_cache:
                return self.location_cache[key], False
            location = Location(
                site=site,
                building=values.get("building") or "",
                room=values.get("room") or "",
                rack=values.get("rack") or "",
                rack_unit=values.get("rack_unit"),
                notes="Created by spreadsheet import.",
            )
            self._mark_created(location)
            self.location_cache[key] = location
            self.location_cache[("str", str(location).casefold())] = location
            return location, True

        raw = values.get("location")
        if not raw:
            return None, False
        key = ("str", raw.casefold())
        if key in self.location_cache:
            return self.location_cache[key], False
        location = Location(site=raw, notes="Imported free-text location.")
        self._mark_created(location)
        self.location_cache[key] = location
        self.location_cache[self._location_key_from_obj(location)] = location
        return location, True

    # -- placeholders ----------------------------------------------------
    def _allocate_placeholder_name(self):
        index = 1
        while f"UNNAMED-{index:03d}".casefold() in self.devices_by_name:
            index += 1
        name = f"UNNAMED-{index:03d}"
        self.placeholder_count += 1
        return name

    def _find_device_reference(self, raw):
        raw = str(raw).strip()
        if not raw:
            return None
        if raw.isdigit():
            device = Device.objects.filter(pk=int(raw)).first()
            if device:
                return device
        return (
            self.devices_by_name.get(raw.casefold())
            or self.devices_by_serial.get(raw.casefold())
            or self.devices_by_asset.get(raw.casefold())
        )

    def _placeholder_device(self, hint, original_reference):
        hint = (hint or "").strip() or "Unidentified device"
        on_match = re.match(r"^(?P<name>.+?)\s+on\s+(?P<where>.+)$", hint, re.IGNORECASE)
        display_name = on_match.group("name").strip() if on_match else hint
        fingerprint = _safe_reference_fingerprint(display_name)
        existing = self.devices_by_fingerprint.get(fingerprint)
        if existing:
            return existing, False
        existing = self.devices_by_name.get(display_name.casefold())
        if existing:
            return existing, False
        name = display_name
        suffix = 2
        while name.casefold() in self.devices_by_name:
            name = f"{display_name} ({suffix})"
            suffix += 1
        device = Device(
            name=name,
            device_type=Device.Type.OTHER,
            status=Device.Status.ACTIVE,
            needs_identifying=True,
            source=Source.IMPORT,
            notes=f"Imported from reference: {original_reference}",
            import_fingerprint=fingerprint,
            last_verified=timezone.localdate(),
        )
        self._mark_created(device)
        self.devices_by_name[name.casefold()] = device
        self.devices_by_fingerprint[fingerprint] = device
        self.placeholder_count += 1
        return device, True

    def _get_or_create_reference_device(self, raw):
        existing = self._find_device_reference(raw)
        if existing:
            return existing, False
        device, created = self._placeholder_device(raw, raw)
        return device, created

    def _port_key(self, device, label):
        device_key = device.pk if device.pk else f"new-{id(device)}"
        return (device_key, label.casefold())

    def _find_port(self, device, label):
        return self.ports_by_key.get(self._port_key(device, label))

    def _create_port_placeholder(self, device, label):
        port = Port(
            device=device,
            label=label,
            port_type=Port.Type.ETHERNET,
            status=Port.Status.UNKNOWN,
        )
        self._mark_created(port)
        self.ports_by_key[self._port_key(device, label)] = port
        return port, True

    # -- device rows -----------------------------------------------------
    def _device_match(self, values):
        serial = (values.get("serial_number") or "").casefold()
        asset = (values.get("asset_id") or "").casefold()
        name = (values.get("name") or "").casefold()
        candidates = {}
        for label, cached, key in (
            ("serial number", self.devices_by_serial, serial),
            ("asset ID", self.devices_by_asset, asset),
            ("name", self.devices_by_name, name),
        ):
            if key:
                found = cached.get(key)
                if found is not None:
                    candidates[(found._meta.label_lower, found.pk)] = found
        if len(candidates) > 1:
            return None, "Serial number, asset ID and name point at different devices."
        if candidates:
            return next(iter(candidates.values())), None
        if not serial and not asset and not name:
            fingerprint = _fingerprint("devices", values)
            found = self.devices_by_fingerprint.get(fingerprint)
            if found:
                return found, None
        return None, None

    def _apply_device(self, row_number, values):
        errors = []
        warnings = []
        details = []
        match, conflict = self._device_match(values)
        if conflict:
            return "error", "", [conflict], [], ""
        identifiers = any(values.get(name) for name in ("name", "serial_number", "asset_id"))
        fingerprint = None
        if match is None and not identifiers:
            fingerprint = _fingerprint("devices", values)

        if match is None:
            name = values.get("name") or ""
            if not name:
                name = self._allocate_placeholder_name()
            device = Device(
                name=name,
                device_type=Device.Type.OTHER,
                status=Device.Status.ACTIVE,
                source=Source.IMPORT,
                needs_identifying=not identifiers,
                last_verified=timezone.localdate(),
            )
            self._mark_created(device)
            action = "new"
        else:
            device = match
            action = "update"

        location, location_created = self._location_from_values(values)
        if location_created:
            warnings.append(f"Location {location!s} will be created.")
        if fingerprint:
            device.import_fingerprint = fingerprint

        assignments = {}
        for name, value in values.items():
            if name in {"location", "site", "building", "room", "rack"}:
                continue
            if name == "name" and not value:
                continue
            assignments[name] = value
        if location is not None:
            assignments["location"] = location
        assignments["source"] = Source.IMPORT
        assignments["last_verified"] = values.get("last_verified") or timezone.localdate()
        if identifiers:
            assignments["needs_identifying"] = False
        if self._is_created(device) and not identifiers:
            assignments["needs_identifying"] = True

        old_name = device.name if device.pk else None
        if "name" in assignments and assignments["name"] != device.name:
            existing = self.devices_by_name.get(str(assignments["name"]).casefold())
            if existing is not None and existing is not device:
                return "error", "", [f"Device name {assignments['name']!r} is already used."], [], ""
        self._assign(device, assignments)
        if old_name and old_name.casefold() != device.name.casefold():
            self.devices_by_name.pop(old_name.casefold(), None)
        self.devices_by_name[device.name.casefold()] = device
        if device.serial_number:
            self.devices_by_serial[str(device.serial_number).casefold()] = device
        if device.asset_id:
            self.devices_by_asset[str(device.asset_id).casefold()] = device
        if device.import_fingerprint:
            self.devices_by_fingerprint[device.import_fingerprint] = device

        if self._is_created(device):
            details.append("Creates device")
        else:
            details.append("Updates existing device")
        if device.needs_identifying:
            action = "placeholder" if action == "new" else action
            warnings.append("Needs identifying after import.")
        return action, device.name, errors, warnings, ", ".join(details)

    def _process_device_row(self, number, row):
        values, issues = parse_fields(row, self.mapping, "devices")
        if issues:
            return PreviewRow(number, "error", "", [], issues)
        if not values:
            return PreviewRow(number, "error", "", [], ["No mapped values in this row."])
        action, subject, errors, warnings, details = self._apply_device(number, values)
        return PreviewRow(number, action, subject, [details] if details else [], errors, warnings)

    # -- port rows -------------------------------------------------------
    def _process_port_row(self, number, row):
        values, issues = parse_fields(row, self.mapping, "ports")
        if issues:
            return PreviewRow(number, "error", "", [], issues)
        if not values.get("label"):
            return PreviewRow(number, "error", "", [], ["A port label or range is required."])
        raw_device = values.get("device")
        if not raw_device:
            return PreviewRow(number, "error", "", [], ["A device reference is required."])
        warnings = []
        if ":" in raw_device:
            raw_device = raw_device.rsplit(":", 1)[0].strip()
        device, device_created = self._get_or_create_reference_device(raw_device)
        if device_created:
            warnings.append(f"Placeholder device {device.name!r} will be created.")

        labels = expand_port_labels(values.get("label"))
        if not labels:
            return PreviewRow(number, "error", "", [], ["Could not read the port label."])
        if len(labels) > 256:
            return PreviewRow(number, "error", "", [], ["That range expands to more than 256 ports."])

        actions = set()
        subjects = []
        errors = []
        for label in labels:
            port = self._find_port(device, label)
            if port is None:
                port = Port(
                    device=device,
                    label=label,
                    port_type=Port.Type.ETHERNET,
                    status=Port.Status.UNKNOWN,
                )
                self._mark_created(port)
                self.ports_by_key[self._port_key(device, label)] = port
                actions.add("new")
            else:
                actions.add("update")
            assignments = {}
            for name in (
                "port_type",
                "speed",
                "poe",
                "access_vlan",
                "tagged_vlans",
                "connector",
                "fibre_type",
                "provider",
                "circuit_id",
                "status",
            ):
                if name in values and values[name] is not None:
                    assignments[name] = values[name]
            self._assign(port, assignments)
            subjects.append(f"{device.name}:{label}")
            if device_created:
                self.placeholder_count += 0  # already counted by placeholder helper
        if len(actions) > 1:
            action = "mixed"
        else:
            action = next(iter(actions)) if actions else "update"
        if device_created and action == "new":
            action = "placeholder"
        expanded = ""
        if len(labels) > 1:
            expanded = ", ".join(labels)
            warnings.append(f"Expanded {values.get('label')!r} into {len(labels)} ports.")
        subject = f"{device.name} · {len(labels)} port{'s' if len(labels) != 1 else ''}"
        details = [", ".join(subjects[:8]) + ("…" if len(subjects) > 8 else "")]
        return PreviewRow(number, action, subject, details, errors, warnings, expanded)

    # -- connection rows -------------------------------------------------
    @staticmethod
    def _split_reference(raw):
        text = str(raw).strip()
        if not text:
            return None
        if text.casefold() in UNKNOWN_VALUES:
            return None
        paren = re.match(
            r"^(?P<device>.+?)\s*\(\s*(?:port\s*)?(?P<port>[^)]+?)\s*\)\s*$",
            text,
            re.IGNORECASE,
        )
        if paren:
            return paren.group("device").strip(), paren.group("port").strip()
        if ":" in text:
            device, port = text.rsplit(":", 1)
            return device.strip(), port.strip()
        port_word = re.match(
            r"^(?P<device>.+?)\s+port\s+(?P<port>[A-Za-z0-9/._-]+)\s*$",
            text,
            re.IGNORECASE,
        )
        if port_word:
            return port_word.group("device").strip(), port_word.group("port").strip()
        return text, "UNKNOWN"

    def _resolve_port_reference(self, raw, side):
        if raw is None or not str(raw).strip():
            return None, False
        parsed = self._split_reference(raw)
        if parsed is None:
            return None, False
        device_hint, label = parsed
        device, device_created = self._get_or_create_reference_device(device_hint)
        port = self._find_port(device, label)
        created = device_created
        if port is None:
            port, port_created = self._create_port_placeholder(device, label)
            created = created or port_created
        return port, created

    @staticmethod
    def _connection_key(port_a_id, port_b_id, raw=False):
        if raw:
            values = [port_a_id, port_b_id]
            values = [value for value in values if value is not None]
            return tuple(sorted(values))
        values = [port_a_id, port_b_id]
        values = [value for value in values if value is not None]
        return tuple(sorted(values, key=str))

    def _port_identity(self, port):
        return port.pk if port.pk else f"new-{id(port)}"

    def _process_connection_row(self, number, row):
        values, issues = parse_fields(row, self.mapping, "connections")
        if issues:
            return PreviewRow(number, "error", "", [], issues)

        raw_a = values.get("port_a")
        if not raw_a and values.get("port_a_device"):
            label = values.get("port_a_label")
            raw_a = (
                f"{values['port_a_device']}:{label}"
                if label
                else values["port_a_device"]
            )
        raw_b = values.get("port_b")
        if not raw_b and values.get("port_b_device"):
            label = values.get("port_b_label")
            raw_b = (
                f"{values['port_b_device']}:{label}"
                if label
                else values["port_b_device"]
            )
        if not raw_a:
            return PreviewRow(number, "error", "", [], ["Port A is required."])

        port_a, created_a = self._resolve_port_reference(raw_a, "A")
        port_b, created_b = self._resolve_port_reference(raw_b, "B")
        if port_a is None:
            return PreviewRow(number, "error", "", [], ["Port A could not be resolved."])
        if port_b is not None and port_a is port_b:
            return PreviewRow(number, "error", "", [], ["A port cannot be connected to itself."])

        key = self._connection_key(self._port_identity(port_a), self._port_identity(port_b))
        connection = self.connection_cache.get(key)
        if connection is None:
            connection = Connection(
                port_a=port_a,
                port_b=port_b,
                source=Source.IMPORT,
            )
            self._mark_created(connection)
            self.connection_cache[key] = connection
            action = "new"
        else:
            action = "update"
        assignments = {}
        for name in ("cable_label", "notes"):
            if name in values and values[name] is not None:
                assignments[name] = values[name]
        self._assign(connection, assignments)

        warnings = []
        if created_a or created_b:
            action = "placeholder"
            warnings.append("Missing port/device references will be created as placeholders.")
        details = [
            f"A: {port_a.device.name}:{port_a.label}",
            f"B: {port_b.device.name + ':' + port_b.label if port_b else 'Unknown'}",
        ]
        return PreviewRow(number, action, str(connection), details, [], warnings)

    # -- orchestration ---------------------------------------------------
    def run(self):
        if self.target not in TARGET_FIELDS:
            return ImportResult(
                rows=[],
                counts={},
                global_errors=["Choose what this sheet contains before previewing."],
                sheet=self.batch.sheet_name,
                target=self.target,
            )
        try:
            data = read_sheet(self.batch, self.batch.sheet_name, self.batch.has_header)
        except ImportFileError as exc:
            return ImportResult(
                rows=[],
                counts={},
                global_errors=[str(exc)],
                sheet=self.batch.sheet_name,
                target=self.target,
            )

        valid_fields = field_names_for(self.target)
        for index, field_name in self.mapping.items():
            if field_name and field_name not in valid_fields:
                self.global_errors.append(f"Column {index} is mapped to an unknown field.")
        mapped_fields = [field for field in self.mapping.values() if field]
        if len(mapped_fields) != len(set(mapped_fields)):
            self.global_errors.append("The same NetMap field is mapped from more than one column.")

        processor = {
            "devices": self._process_device_row,
            "ports": self._process_port_row,
            "connections": self._process_connection_row,
        }[self.target]
        for number, row in data.rows:
            try:
                preview = processor(number, row)
            except Exception as exc:  # Keep the preview useful even for a bad row.
                preview = PreviewRow(number, "error", "", [], [f"Unexpected row error: {exc}"])
            self.rows.append(preview)

        self._validate_objects()
        counts = {
            "new": sum(1 for row in self.rows if row.action == "new"),
            "update": sum(1 for row in self.rows if row.action == "update"),
            "placeholder": sum(1 for row in self.rows if row.action == "placeholder"),
            "mixed": sum(1 for row in self.rows if row.action == "mixed"),
            "error": sum(1 for row in self.rows if row.action == "error" or row.errors),
        }
        return ImportResult(
            rows=self.rows,
            counts=counts,
            created_count=len(self.created),
            updated_count=len(self.updated),
            placeholder_count=self.placeholder_count,
            global_errors=self.global_errors,
            total_rows=len(data.rows),
            sheet=self.batch.sheet_name,
            target=self.target,
        )

    def _validate_objects(self):
        for obj in self.created:
            exclude = []
            for fk_name in ("location", "device", "port_a", "port_b"):
                if not hasattr(obj, fk_name):
                    continue
                related = getattr(obj, fk_name, None)
                if related is not None and getattr(related, "pk", None) is None:
                    exclude.append(fk_name)
            try:
                obj.full_clean(exclude=exclude)
            except ValidationError as exc:
                label = f"New {obj._meta.verbose_name} {obj!s}: "
                for field_name, messages in exc.message_dict.items():
                    for message in messages:
                        self.global_errors.append(label + f"{field_name}: {message}")
        for obj in self.updated.values():
            try:
                obj.full_clean()
            except ValidationError as exc:
                label = f"Updated {obj._meta.verbose_name} {obj!s}: "
                for field_name, messages in exc.message_dict.items():
                    for message in messages:
                        self.global_errors.append(label + f"{field_name}: {message}")

    def commit(self):
        result = self.run()
        if result.has_errors:
            raise ImportCommitError("Fix the preview errors before saving.")
        order = [Location, Device, Port, IPAddress, Connection]
        with transaction.atomic():
            for model in order:
                for obj in self.created:
                    if isinstance(obj, model):
                        _sync_relation_ids(obj)
                        obj.full_clean()
                        obj.save()
            for obj in self.updated.values():
                _sync_relation_ids(obj)
                obj.full_clean()
                obj.save()
        undo_data = {
            "created": [
                {"model": obj._meta.label_lower, "pk": obj.pk}
                for obj in self.created
                if obj.pk
            ],
            "updated": [
                {
                    "model": obj._meta.label_lower,
                    "pk": obj.pk,
                    "fields": self.snapshots[(obj._meta.label_lower, obj.pk)],
                }
                for obj in self.updated.values()
            ],
        }
        return result, undo_data


# ---------------------------------------------------------------------------
# Snapshots and undo
# ---------------------------------------------------------------------------

SNAPSHOT_EXCLUDE = {"created_at", "updated_at"}


def _serialize_value(value):
    if isinstance(value, datetime):
        return {"__type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"__type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"__type": "time", "value": value.isoformat()}
    if isinstance(value, Decimal):
        return {"__type": "decimal", "value": str(value)}
    return value


def _sync_relation_ids(obj):
    """After parents are saved, copy cached related pks onto FK attnames.

    Assigning an unsaved related instance does not populate ``field_id`` until
    the object itself is saved. full_clean() needs the id when the target is now
    available.
    """
    for field in obj._meta.concrete_fields:
        if not field.is_relation:
            continue
        related = getattr(obj, field.name, None)
        if related is not None and getattr(related, "pk", None):
            setattr(obj, field.attname, related.pk)


def serialize_instance(obj):
    data = {}
    for field in obj._meta.concrete_fields:
        if field.name in SNAPSHOT_EXCLUDE or field.primary_key:
            continue
        data[field.attname] = _serialize_value(getattr(obj, field.attname))
    return data


def _deserialize_value(value):
    if isinstance(value, dict) and "__type" in value:
        if value["__type"] == "date":
            return date.fromisoformat(value["value"])
        if value["__type"] == "datetime":
            return datetime.fromisoformat(value["value"])
        if value["__type"] == "time":
            return time.fromisoformat(value["value"])
        if value["__type"] == "decimal":
            return Decimal(value["value"])
    return value


def undo_import(batch: ImportBatch):
    if batch.status != ImportBatch.Status.IMPORTED:
        raise ImportUndoError("Only a completed import can be undone.")
    undo_data = batch.undo_data or {}
    order = ["connection", "port", "ipaddress", "device", "location"]
    model_models = {
        model._meta.model_name: model
        for model in (Connection, Port, IPAddress, Device, Location)
    }
    with transaction.atomic():
        created = sorted(
            undo_data.get("created", []),
            key=lambda item: order.index(item["model"].split(".")[-1])
            if item["model"].split(".")[-1] in order
            else len(order),
        )
        for item in created:
            model_name = item["model"].split(".")[-1]
            model = model_models.get(model_name)
            if model is None:
                continue
            obj = model.objects.filter(pk=item["pk"]).first()
            if obj is None:
                continue
            try:
                obj.delete()
            except Exception as exc:
                raise ImportUndoError(
                    f"Could not undo: {obj} is now referenced elsewhere ({exc})."
                ) from exc

        for item in undo_data.get("updated", []):
            model_name = item["model"].split(".")[-1]
            model = model_models.get(model_name)
            if model is None:
                continue
            obj = model.objects.filter(pk=item["pk"]).first()
            if obj is None:
                continue
            for attname, value in item.get("fields", {}).items():
                setattr(obj, attname, _deserialize_value(value))
            try:
                obj.save()
            except Exception as exc:
                raise ImportUndoError(
                    f"Could not restore {obj}: {exc}"
                ) from exc
    batch.status = ImportBatch.Status.UNDONE
    batch.undone_at = timezone.now()
    batch.save(update_fields=["status", "undone_at"])
